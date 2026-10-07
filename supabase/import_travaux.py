import argparse
import json
from datetime import date, datetime
from pathlib import Path

import openpyxl
import import_baux as api


def scalar(value):
    return value.isoformat()[:10] if isinstance(value, (date, datetime)) else value


def extract(path):
    workbook = openpyxl.load_workbook(path, data_only=True)
    register = workbook["2_Registre travaux"]
    headers = [api.text(register.cell(4, c).value) for c in range(1, 38)]
    ppi = {}
    for row in workbook["5_PPI 2026-2030"].iter_rows(min_row=6, values_only=True):
        if isinstance(row[1], (int, float)):
            ppi[int(row[1])] = {str(2026 + i): api.number(row[5 + i]) for i in range(5)}
    studies = {}
    for row in workbook["3_Études à mener"].iter_rows(min_row=5, values_only=True):
        if isinstance(row[1], (int, float)):
            studies[int(row[1])] = [scalar(v) for v in row[2:]]
    rows = []
    for values in register.iter_rows(min_row=5, values_only=True):
        if not isinstance(values[0], (int, float)) or not values[5]:
            continue
        n = int(values[0])
        details = {
            "version_travaux": 1, "source_key": f"ruvelon_travaux_v2_{n}",
            "numero": n, "annee": int(values[34]) if values[34] else None,
            "localisation": " / ".join(api.text(v) for v in values[3:5] if v),
            "categorie": api.text(values[2]), "devis_ht": api.number(values[30]) if values[30] is not None else None,
            "note_matt": api.text(values[9]), "imputation": api.text(values[13]),
            "notes": api.text(values[36]), "ppi": ppi.get(n, {}),
            "etude": studies.get(n), "source": {k: scalar(v) for k, v in zip(headers, values)},
        }
        rows.append({
            "titre": api.text(values[5]), "description": api.text(values[6]),
            "priorite": api.text(values[10]) or "À définir", "statut": api.text(values[18]) or "0 - À chiffrer",
            "entreprise": api.text(values[19]), "montant_estime": api.number(values[29]),
            "montant_reel": api.number(values[31]),
            "date_prevue": None, "date_realisation": scalar(values[27]),
            "observations": json.dumps(details, ensure_ascii=False),
        })
    return rows


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook")
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()
    rows = extract(args.workbook)
    assert len(rows) == 36
    assert round(sum(r["montant_estime"] for r in rows), 2) == 116990
    created = skipped = 0
    if args.apply:
        token = api.request("/auth/v1/token?grant_type=password", method="POST", body={"email": api.ADMIN_EMAIL, "password": api.ADMIN_PASSWORD})["access_token"]
        existing = api.request("/rest/v1/travaux?select=*", token=token) or []
        keys = set()
        for row in existing:
            try:
                keys.add(json.loads(row.get("observations") or "{}").get("source_key"))
            except (ValueError, AttributeError):
                pass
        for row in rows:
            key = json.loads(row["observations"])["source_key"]
            if key in keys or any(api.normalize(x["titre"]) == api.normalize(row["titre"]) for x in existing):
                skipped += 1
                continue
            api.request("/rest/v1/travaux", method="POST", body=row, token=token, prefer="return=minimal")
            keys.add(key)
            created += 1
        verified = api.request("/rest/v1/travaux?select=*", token=token) or []
        assert len(verified) == len(existing) + created
    print(json.dumps({"operations": len(rows), "estimation_ht": round(sum(r["montant_estime"] for r in rows), 2), "devis_ht": round(sum(json.loads(r["observations"])["devis_ht"] or 0 for r in rows), 2), "reel_ht": round(sum(r["montant_reel"] for r in rows), 2), "creees": created, "existantes_conservees": skipped}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
