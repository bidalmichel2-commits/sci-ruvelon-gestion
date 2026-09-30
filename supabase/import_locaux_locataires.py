import json
import os
import re
import sys
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
from collections import defaultdict
from datetime import date, datetime

import openpyxl


SUPABASE_URL = os.environ.get("SCI_SUPABASE_URL", "").rstrip("/")
SUPABASE_KEY = os.environ.get("SCI_SUPABASE_KEY", "")
ADMIN_EMAIL = os.environ.get("SCI_ADMIN_EMAIL", "")
ADMIN_PASSWORD = os.environ.get("SCI_ADMIN_PASSWORD", "")


def request(path, method="GET", body=None, token=None, prefer=None):
    headers = {"apikey": SUPABASE_KEY, "Content-Type": "application/json"}
    if token:
        headers["Authorization"] = "Bearer " + token
    if prefer:
        headers["Prefer"] = prefer
    data = None if body is None else json.dumps(body, ensure_ascii=False).encode("utf-8")
    req = urllib.request.Request(SUPABASE_URL + path, data=data, headers=headers, method=method)
    try:
        with urllib.request.urlopen(req, timeout=30) as response:
            raw = response.read()
            return json.loads(raw.decode("utf-8")) if raw else None
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"Supabase {exc.code}: {detail}") from exc


def text(value):
    return "" if value is None else str(value).strip()


def number(value):
    try:
        return round(float(value or 0), 2)
    except (TypeError, ValueError):
        return 0


def display_date(value):
    if isinstance(value, (datetime, date)):
        return value.strftime("%d/%m/%Y")
    return text(value)


def normalize(value):
    value = unicodedata.normalize("NFKD", text(value)).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"\s+", " ", value).strip().lower()


def split_codes(value):
    return [part.strip() for part in text(value).split("+") if part.strip()]


def split_person(label):
    activity = ""
    match = re.match(r"^(.*?)\s*\((.*?)\)\s*$", label)
    if match:
        label, activity = match.group(1).strip(), match.group(2).strip()
    parts = label.split()
    return (parts[0] if parts else "", " ".join(parts[1:]), activity)


def source_note(row, local_codes):
    total = number(row[6])
    charges = number(row[7])
    hors_charges = number(row[8])
    caution = number(row[9])
    parts = [
        "Locaux : " + "+".join(local_codes),
        "Entrée : " + display_date(row[3]),
        "Sortie prévue : " + display_date(row[5]),
        f"Loyer hors charges : {hors_charges:.2f} €",
        f"Charges : {charges:.2f} €",
        f"Total avec charges : {total:.2f} €",
    ]
    if caution:
        parts.append(f"Dépôt de garantie : {caution:.2f} €")
    if text(row[10]):
        parts.append("État du local : " + text(row[10]))
    if text(row[13]):
        parts.append(text(row[13]))
    return ". ".join(parts) + "."


def main(workbook_path):
    required = [SUPABASE_URL, SUPABASE_KEY, ADMIN_EMAIL, ADMIN_PASSWORD]
    if not all(required):
        raise RuntimeError("Configuration Supabase incomplète.")

    workbook = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    local_sheet = workbook["2. Locaux"]
    tenant_sheet = workbook["3. Locataires"]

    company_names = {
        "LMT BISCUITERIE (LOC MARIA)",
        "EUR WORD SOCCER MANAGEMENT",
        "FORMACOOPP RELACOM",
        "ORYX",
    }

    tenants = []
    occupants = defaultdict(list)
    for row in tenant_sheet.iter_rows(min_row=2, values_only=True):
        label = text(row[0])
        if not label or label.upper() == "TOTAL" or normalize(row[2]) != "actif":
            continue
        codes = split_codes(row[1])
        if label in company_names:
            nom, prenom, raison_sociale, activity = "", "", label, ""
            tenant_type = "Société"
        else:
            nom, prenom, activity = split_person(label)
            raison_sociale = ""
            tenant_type = "Particulier"
        email = text(row[12])
        if email == "coralie,boissiere@gmail,com":
            email = "coralie.boissiere@gmail.com"
        tenant = {
            "type": tenant_type,
            "nom": nom,
            "prenom": prenom,
            "raison_sociale": raison_sociale,
            "activite": activity,
            "telephone": text(row[11]),
            "email": email,
            "statut": "Actif",
            "observations": source_note(row, codes),
            "_label": label,
            "_codes": codes,
            "_loyer_hc": number(row[8]),
            "_charges": number(row[7]),
            "_caution": number(row[9]),
        }
        tenants.append(tenant)
        for code in codes:
            occupants[code].append(tenant)

    locaux = []
    for row in local_sheet.iter_rows(min_row=2, values_only=True):
        code = text(row[0])
        if not code or code.upper() == "TOTAL":
            continue
        kind = text(row[4])
        current = occupants.get(code, [])
        is_common = normalize(kind) == "commun"
        if is_common:
            status = "Partie commune"
        elif current:
            status = "Loue"
        else:
            status = "Disponible"

        amounts_are_local = bool(current) and all(len(t["_codes"]) == 1 for t in current)
        rent = round(sum(t["_loyer_hc"] for t in current), 2) if amounts_are_local else 0
        charges = round(sum(t["_charges"] for t in current), 2) if amounts_are_local else 0
        deposit = round(sum(t["_caution"] for t in current), 2) if amounts_are_local else 0

        notes = []
        if current:
            notes.append("Occupant(s) : " + ", ".join(t["_label"] for t in current))
            if not amounts_are_local:
                notes.append("Loyer global non ventilé entre plusieurs locaux dans le tableau source")
        elif not is_common:
            notes.append("Vacant selon le tableau source")
        if text(row[7]):
            notes.append(text(row[7]))
        if text(row[5]):
            meter = "Compteur énergie : " + text(row[5])
            if row[6] is not None:
                meter += f" ({float(row[6]) * 100:.2f} %)"
            notes.append(meter)
        if row[3] is not None and number(row[3]) != number(row[2]):
            notes.append(f"Surface avec communs : {number(row[3]):.2f} m²")
        notes.append("Source : Liste location 2026.xlsx")

        locaux.append({
            "code_local": code,
            "designation": text(row[7]) if is_common and text(row[7]) else "Local " + code,
            "niveau": text(row[1]),
            "surface_m2": number(row[2]),
            "statut": status,
            "loyer_ht": rent,
            "tva_loyer": 0,
            "loyer_ttc": rent,
            "charges_mensuelles": charges,
            "depot_garantie": deposit,
            "observations": ". ".join(notes) + ".",
            "actif": True,
        })

    auth = request(
        "/auth/v1/token?grant_type=password",
        method="POST",
        body={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD},
    )
    token = auth["access_token"]

    existing_locals = request("/rest/v1/locaux?select=code_local", token=token) or []
    existing_codes = {normalize(row.get("code_local")) for row in existing_locals}
    new_locals = [row for row in locaux if normalize(row["code_local"]) not in existing_codes]
    if new_locals:
        request("/rest/v1/locaux", method="POST", body=new_locals, token=token, prefer="return=representation")

    existing_tenants = request(
        "/rest/v1/locataires?select=nom,prenom,raison_sociale,email", token=token
    ) or []
    existing_keys = set()
    for row in existing_tenants:
        key = normalize(row.get("email")) or normalize(row.get("raison_sociale")) or normalize(
            f"{row.get('nom', '')} {row.get('prenom', '')}"
        )
        existing_keys.add(key)

    new_tenants = []
    for tenant in tenants:
        clean = {key: value for key, value in tenant.items() if not key.startswith("_")}
        key = normalize(clean.get("email")) or normalize(clean.get("raison_sociale")) or normalize(
            f"{clean.get('nom', '')} {clean.get('prenom', '')}"
        )
        if key not in existing_keys:
            new_tenants.append(clean)
            existing_keys.add(key)
    if new_tenants:
        request(
            "/rest/v1/locataires",
            method="POST",
            body=new_tenants,
            token=token,
            prefer="return=representation",
        )

    final_locals = request(
        "/rest/v1/locaux?select=code_local,statut&order=code_local", token=token
    ) or []
    final_tenants = request(
        "/rest/v1/locataires?select=nom,prenom,raison_sociale,statut&order=created_at", token=token
    ) or []
    summary = {
        "locaux_ajoutes": len(new_locals),
        "locataires_ajoutes": len(new_tenants),
        "locaux_total": len(final_locals),
        "locataires_total": len(final_tenants),
        "disponibles": sorted(row["code_local"] for row in final_locals if row["statut"] == "Disponible"),
        "parties_communes": sorted(
            row["code_local"] for row in final_locals if row["statut"] == "Partie commune"
        ),
    }
    print(json.dumps(summary, ensure_ascii=False))


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("Usage: import_locaux_locataires.py <classeur.xlsx>")
    main(sys.argv[1])
