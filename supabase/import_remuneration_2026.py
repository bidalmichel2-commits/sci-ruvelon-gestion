import argparse
import json
import os
import re
import urllib.error
import urllib.parse
import urllib.request
from datetime import date, datetime

import openpyxl


SUPABASE_URL = os.environ.get("SCI_SUPABASE_URL", "").rstrip("/")
SUPABASE_KEY = os.environ.get("SCI_SUPABASE_KEY", "")
ADMIN_EMAIL = os.environ.get("SCI_ADMIN_EMAIL", "")
ADMIN_PASSWORD = os.environ.get("SCI_ADMIN_PASSWORD", "")

MONTH_NUMBERS = {
    "JANVIER": 1,
    "FÉVRIER": 2,
    "MARS": 3,
    "AVRIL": 4,
    "MAI": 5,
    "JUIN": 6,
    "JUILLET": 7,
    "AOÛT": 8,
    "SEPTEMBRE": 9,
    "OCTOBRE": 10,
    "NOVEMBRE": 11,
    "DÉCEMBRE": 12,
}


def number(value):
    try:
        return round(float(value or 0), 2)
    except (TypeError, ValueError):
        return 0.0


def text(value):
    return "" if value is None else str(value).strip()


def display_date(value):
    if isinstance(value, (date, datetime)):
        return value.strftime("%d/%m/%Y")
    return text(value).split(" ")[0]


def iso_date(value):
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    raw = text(value).split(" ")[0]
    if not raw:
        return ""
    try:
        return datetime.fromisoformat(raw).strftime("%Y-%m-%d")
    except ValueError:
        return raw


def request(path, method="GET", body=None, token=None, prefer=None):
    headers = {"apikey": SUPABASE_KEY, "Content-Type": "application/json"}
    if token:
        headers["Authorization"] = "Bearer " + token
    if prefer:
        headers["Prefer"] = prefer
    payload = None if body is None else json.dumps(body, ensure_ascii=False).encode("utf-8")
    req = urllib.request.Request(SUPABASE_URL + path, data=payload, headers=headers, method=method)
    try:
        with urllib.request.urlopen(req, timeout=30) as response:
            raw = response.read()
            return json.loads(raw.decode("utf-8")) if raw else None
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"Supabase {exc.code}: {detail}") from exc


def encaissement_lines(sheet):
    rows = []
    for row in range(9, 20):
        local = text(sheet.cell(row, 1).value)
        tenant = text(sheet.cell(row, 2).value)
        rent = number(sheet.cell(row, 3).value)
        received = number(sheet.cell(row, 4).value)
        observation = text(sheet.cell(row, 5).value)
        if not any((tenant, rent, received, observation)):
            continue
        rows.append({
            "local": local,
            "locataire": tenant,
            "loyer_ht": rent,
            "encaissement_ht": received,
            "observations": observation,
        })
    return rows


def intervention_lines(sheet):
    rows = []
    for row in range(24, 44):
        activity = text(sheet.cell(row, 3).value)
        local = text(sheet.cell(row, 4).value)
        hours = number(sheet.cell(row, 5).value)
        km = number(sheet.cell(row, 6).value)
        other = number(sheet.cell(row, 8).value)
        observation = text(sheet.cell(row, 9).value)
        if not any((activity, local, hours, km, other, observation)):
            continue
        rows.append({
            "date": iso_date(sheet.cell(row, 1).value),
            "activite": activity,
            "local": local,
            "heures": hours,
            "km": km,
            "autres_frais": other,
            "observations": observation,
        })
    return rows


def prime_situation(reference, prime):
    if not reference:
        return "Montant manuel"
    ratio = round(prime / reference, 4)
    if abs(ratio - 1) <= 0.001:
        return "Sans interruption"
    if abs(ratio - 0.5) <= 0.001:
        return "Interruption moins 1 mois"
    if abs(ratio - 0.25) <= 0.001:
        return "Interruption plus 1 mois"
    return "Montant manuel"


def prime_lines(sheet):
    rows = []
    for row in range(48, 53):
        values = [sheet.cell(row, col).value for col in range(1, 7)]
        if not any(value not in (None, "") for value in values):
            continue
        reference = number(values[3])
        prime = number(values[4])
        rows.append({
            "date": iso_date(values[0]),
            "local": text(values[1]),
            "type": text(values[2]),
            "loyer_reference": reference,
            "situation": prime_situation(reference, prime),
            "prime": prime,
            "observations": text(values[5]),
        })
    return rows


def prime_notes(sheet):
    notes = []
    references = []
    for row in range(48, 53):
        values = [sheet.cell(row, col).value for col in range(1, 7)]
        if not any(value not in (None, "") for value in values):
            continue
        reference = number(values[3])
        prime = number(values[4])
        if reference:
            references.append(reference)
        parts = [display_date(values[0]), text(values[1]), text(values[2])]
        if reference:
            parts.append(f"loyer reference {reference:.2f} EUR")
        if prime:
            parts.append(f"prime {prime:.2f} EUR")
        if text(values[5]):
            parts.append(text(values[5]))
        notes.append(" - ".join(part for part in parts if part))
    return notes, round(sum(references), 2)


def workbook_year(workbook):
    title = text(workbook["RÉCAPITULATIF ANNUEL"].cell(1, 1).value)
    match = re.search(r"\b(20\d{2})\b", title)
    if not match:
        raise RuntimeError("Annee introuvable dans le titre du recapitulatif annuel.")
    return int(match.group(1))


def extract_rows(workbook_path):
    workbook = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    recap = workbook["RÉCAPITULATIF ANNUEL"]
    year = workbook_year(workbook)
    include_other = year != 2026
    rows = []
    checks = []

    for row_number in range(5, 17):
        month_name = text(recap.cell(row_number, 1).value).upper()
        month_number = MONTH_NUMBERS.get(month_name)
        if not month_number:
            continue
        sheet = workbook[month_name]
        encaissements = number(sheet.cell(20, 4).value)
        prime = number(sheet.cell(53, 5).value)
        km = number(sheet.cell(44, 6).value)
        other = number(sheet.cell(44, 8).value)
        hours = number(sheet.cell(44, 5).value)
        encaissement_details = encaissement_lines(sheet)
        intervention_details = intervention_lines(sheet)
        prime_details = prime_lines(sheet)

        if not any((encaissements, prime, km, other, hours, encaissement_details, intervention_details, prime_details)):
            continue

        primes, reference_total = prime_notes(sheet)
        details = {
            "version": 1,
            "encaissements": encaissement_details,
            "interventions": intervention_details,
            "primes": prime_details,
            "notes": (
                f"Import depuis {os.path.basename(workbook_path)}. "
                + ("Les autres frais sont inclus dans le total du compte courant."
                   if include_other else
                   "Les autres frais 2026 sont indiques mais exclus du total du compte courant.")
            ),
        }
        calculated_total = round(
            encaissements * 0.07 * 0.70 + km * 0.636 + prime + (other if include_other else 0),
            2,
        )
        # En 2026, les autres frais sont seulement informatifs : le recapitulatif
        # annuel contient donc le total de reference hors autres frais.
        source_total = number(
            sheet.cell(61, 4).value if include_other else recap.cell(row_number, 8).value
        )
        checks.append({
            "mois": month_name.title(),
            "source": source_total,
            "calcule": calculated_total,
            "ecart": round(source_total - calculated_total, 2),
        })

        rows.append({
            "annee": year,
            "mois": f"{year}-{month_number:02d}-01",
            "encaissements_ht": encaissements,
            "heures": hours,
            "km": km,
            "bareme_km": 0.636,
            "loyer_reference_prime": reference_total,
            "situation_prime": "Montant manuel" if prime else "Aucune prime",
            "taux_prime": 0,
            "prime_edl": prime,
            "prime_responsabilite": 0,
            "autres_frais": other,
            "autres_frais_inclus_total": include_other,
            "note_prime": " | ".join(primes),
            "observations": json.dumps(details, ensure_ascii=False),
        })

    return rows, checks, year


def authenticate():
    if not all((SUPABASE_URL, SUPABASE_KEY, ADMIN_EMAIL, ADMIN_PASSWORD)):
        raise RuntimeError("Configuration Supabase incomplete.")
    auth = request(
        "/auth/v1/token?grant_type=password",
        method="POST",
        body={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD},
    )
    return auth["access_token"]


def import_rows(rows, year):
    token = authenticate()
    existing = request(
        f"/rest/v1/remuneration_gerant?select=id,mois&annee=eq.{year}",
        token=token,
    ) or []
    existing_months = {str(item["mois"])[:10] for item in existing}
    result = request(
        "/rest/v1/remuneration_gerant?on_conflict=mois",
        method="POST",
        body=rows,
        token=token,
        prefer="resolution=merge-duplicates,return=representation",
    ) or []
    return {
        "lignes_traitees": len(result),
        "lignes_creees": sum(1 for row in rows if row["mois"] not in existing_months),
        "lignes_mises_a_jour": sum(1 for row in rows if row["mois"] in existing_months),
    }


def main():
    parser = argparse.ArgumentParser(description="Import annuel du suivi de remuneration")
    parser.add_argument("workbook")
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    rows, checks, year = extract_rows(args.workbook)
    if any(abs(check["ecart"]) > 0.02 for check in checks):
        raise RuntimeError("Un total mensuel ne correspond pas aux regles de calcul.")

    summary = {
        "annee": year,
        "mois_extraits": len(rows),
        "total_encaissements_ht": round(sum(row["encaissements_ht"] for row in rows), 2),
        "total_heures": round(sum(row["heures"] for row in rows), 2),
        "total_km": round(sum(row["km"] for row in rows), 2),
        "total_primes": round(sum(row["prime_edl"] for row in rows), 2),
        "autres_frais": round(sum(row["autres_frais"] for row in rows), 2),
        "autres_frais_inclus_total": year != 2026,
        "total_compte_courant": round(sum(check["calcule"] for check in checks), 2),
        "controles": checks,
    }
    if not args.dry_run:
        summary.update(import_rows(rows, year))
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
